from celery import shared_task
from django.core.files.base import ContentFile
from django.db import transaction
from docx import Document as DocxDocument
import io
from .models import Document, Section, SectionContent, ExportJob
from .services import normalize_text, log_action

@shared_task
def parse_docx(document_id: str):
    doc = Document.objects.get(id=document_id)
    doc.parse_status = Document.ParseStatus.RUNNING
    doc.save(update_fields=["parse_status"])

    try:
        d = DocxDocument(doc.source_file.path)
        # Best-effort parse: Heading styles become sections
        current_section = None
        order = 0
        with transaction.atomic():
            Section.objects.filter(document=doc).delete()

            for para in d.paragraphs:
                text = (para.text or "").strip()
                if not text:
                    continue
                style = (para.style.name if para.style else "") or ""
                is_heading = style.lower().startswith("heading")
                if is_heading:
                    depth = 1
                    # try parse heading level
                    try:
                        depth = int(style.split()[-1])
                    except Exception:
                        depth = 1
                    current_section = Section.objects.create(
                        document=doc,
                        parent=None,  # v1: we keep flat on parse; user can promote/demote later
                        order_index=order,
                        heading=text,
                        depth=depth,
                        old_content_raw="",
                        old_content_normalized="",
                    )
                    SectionContent.objects.create(section=current_section, schema_version=1, blocks_json={"version":1,"blocks":[]})
                    order += 1
                else:
                    if current_section is None:
                        # create a default section
                        current_section = Section.objects.create(
                            document=doc,
                            parent=None,
                            order_index=order,
                            heading="(Untitled)",
                            depth=1,
                            old_content_raw="",
                            old_content_normalized="",
                        )
                        SectionContent.objects.create(section=current_section, schema_version=1, blocks_json={"version":1,"blocks":[]})
                        order += 1
                    current_section.old_content_raw += text + "\n\n"

            # normalize
            for s in Section.objects.filter(document=doc):
                s.old_content_normalized = normalize_text(s.old_content_raw)
                s.save(update_fields=["old_content_normalized"])

        doc.parse_status = Document.ParseStatus.DONE
        doc.save(update_fields=["parse_status"])
        log_action("document", doc.id, "parse_done", {"sections": Section.objects.filter(document=doc).count()})
    except Exception as e:
        doc.parse_status = Document.ParseStatus.FAILED
        doc.save(update_fields=["parse_status"])
        log_action("document", doc.id, "parse_failed", {"error": str(e)})
        raise

@shared_task
def export_docx(export_job_id: str):
    job = ExportJob.objects.get(id=export_job_id)
    job.status = ExportJob.Status.RUNNING
    job.save(update_fields=["status"])

    try:
        # Minimal v1 export stub: create DOCX with headings + plain text flatten.
        # Full template mapping is implemented as next-step: map styles from Template.mapping_json.
        from docx import Document as DocxDocument
        out = DocxDocument()
        doc = job.document
        sections = Section.objects.filter(document=doc).order_by("order_index")
        for s in sections:
            out.add_heading(s.heading or "Section", level=min(max(s.depth,1),4))
            # flatten blocks for now (placeholder)
            if hasattr(s, "content") and s.content.blocks_json:
                bj = s.content.blocks_json
                for b in bj.get("blocks", []):
                    t = b.get("type")
                    if t == "heading":
                        out.add_heading(b.get("text",""), level=min(max(int(b.get("level",2)),1),4))
                    elif t == "paragraph":
                        out.add_paragraph(b.get("text",""))
                    elif t == "list":
                        for item in b.get("items", []):
                            out.add_paragraph(item, style="List Bullet" if b.get("style")=="bulleted" else "List Number")
                    elif t == "image":
                        out.add_paragraph(f"[Image: {b.get('assetId')}] {b.get('caption','')}".strip())
            else:
                out.add_paragraph("")

        bio = io.BytesIO()
        out.save(bio)
        bio.seek(0)
        filename = f"export_{doc.id}.docx"
        job.output_file.save(filename, ContentFile(bio.read()), save=True)
        job.status = ExportJob.Status.DONE
        job.save(update_fields=["status", "output_file"])
        log_action("export", job.id, "export_done", {"file": job.output_file.name})
    except Exception as e:
        job.status = ExportJob.Status.FAILED
        job.error = str(e)
        job.save(update_fields=["status", "error"])
        log_action("export", job.id, "export_failed", {"error": str(e)})
        raise
